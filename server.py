from flask import Flask, jsonify, request, send_from_directory, session, send_file
from functools import wraps
import sqlite3, json, os, secrets, uuid, io
from datetime import datetime
import openpyxl
import urllib.request, urllib.parse, base64

BASE    = os.path.dirname(os.path.abspath(__file__))
# Persistent data directory: Railway Volume at /data, fallback to repo dir locally
DATA_DIR = '/data' if os.path.isdir('/data') else BASE
DB      = os.path.join(DATA_DIR, 'serviceprotokoll.db')
UPLOADS = os.path.join(DATA_DIR, 'uploads')
os.makedirs(UPLOADS, exist_ok=True)
app  = Flask(__name__, static_folder=BASE)
_key_file = os.path.join(DATA_DIR, '.secret_key')
if os.environ.get('SECRET_KEY'):
    app.secret_key = os.environ['SECRET_KEY']
elif os.path.exists(_key_file):
    app.secret_key = open(_key_file).read().strip()
else:
    app.secret_key = secrets.token_hex(32)
    open(_key_file, 'w').write(app.secret_key)

# ── Database ────────────────────────────────────────────────

def get_db():
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    return conn

def init_db():
    with get_db() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS machines (
            nr         TEXT PRIMARY KEY,
            kund       TEXT DEFAULT '',
            anlaggning TEXT DEFAULT '',
            fabrikat   TEXT DEFAULT '',
            modell     TEXT DEFAULT '',
            inkopar    TEXT DEFAULT '',
            notering   TEXT DEFAULT '',
            adress     TEXT DEFAULT '',
            stad       TEXT DEFAULT '',
            kontakt    TEXT DEFAULT '',
            telefon    TEXT DEFAULT '',
            tillvnr    TEXT DEFAULT '',
            arsmodell  TEXT DEFAULT '',
            beskrivning TEXT DEFAULT '',
            agare      TEXT DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS protocols (
            id               INTEGER PRIMARY KEY AUTOINCREMENT,
            maskin_nr        TEXT NOT NULL,
            datum            TEXT,
            type             TEXT,
            kund             TEXT,
            anlaggning       TEXT,
            tekniker         TEXT,
            modell           TEXT,
            items_json       TEXT,
            saved_at         TEXT,
            resolved_json    TEXT DEFAULT '{}',
            workflow_status  TEXT DEFAULT 'inkommen',
            workflow_log     TEXT DEFAULT '[]',
            machine_type     TEXT DEFAULT ''
        );
        CREATE INDEX IF NOT EXISTS idx_protocols_maskin ON protocols(maskin_nr);
        CREATE TABLE IF NOT EXISTS machine_notes (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            maskin_nr TEXT NOT NULL,
            saved_at  TEXT,
            author    TEXT,
            note      TEXT
        );
        CREATE TABLE IF NOT EXISTS offers (
            id           INTEGER PRIMARY KEY AUTOINCREMENT,
            protocol_id  INTEGER,
            maskin_nr    TEXT,
            kund         TEXT,
            anlaggning   TEXT,
            modell       TEXT,
            datum        TEXT,
            tekniker     TEXT,
            items_json   TEXT,
            saved_at     TEXT,
            status       TEXT DEFAULT 'utkast'
        );
        CREATE TABLE IF NOT EXISTS users (
            id       INTEGER PRIMARY KEY AUTOINCREMENT,
            name     TEXT NOT NULL,
            pin      TEXT NOT NULL,
            role     TEXT DEFAULT 'tekniker'
        );
        """)
    # Migration: lägg till kolumner om de saknas
    with get_db() as db:
        cols = [r[1] for r in db.execute("PRAGMA table_info(machines)").fetchall()]
        if 'mot_kaj' not in cols:
            db.execute("ALTER TABLE machines ADD COLUMN mot_kaj TEXT DEFAULT ''")
        if 'hydraulfilter' not in cols:
            db.execute("ALTER TABLE machines ADD COLUMN hydraulfilter TEXT DEFAULT ''")
        if 'hydraulolja' not in cols:
            db.execute("ALTER TABLE machines ADD COLUMN hydraulolja TEXT DEFAULT ''")
        if 'hydraulslangar' not in cols:
            db.execute("ALTER TABLE machines ADD COLUMN hydraulslangar TEXT DEFAULT ''")
        if 'instruktion' not in cols:
            db.execute("ALTER TABLE machines ADD COLUMN instruktion TEXT DEFAULT ''")
        if 'gasdampare' not in cols:
            db.execute("ALTER TABLE machines ADD COLUMN gasdampare TEXT DEFAULT ''")
        if 'typskylt' not in cols:
            db.execute("ALTER TABLE machines ADD COLUMN typskylt TEXT DEFAULT ''")
    # Skapa hydrauliklogg-tabell om den saknas
    with get_db() as db:
        db.execute("""
        CREATE TABLE IF NOT EXISTS hydrauliklogg (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            maskin_nr TEXT NOT NULL,
            datum     TEXT NOT NULL,
            filter    INTEGER DEFAULT 0,
            olja      INTEGER DEFAULT 0,
            slangar   INTEGER DEFAULT 0,
            notering  TEXT DEFAULT '',
            saved_at  TEXT
        )""")
        db.execute("CREATE INDEX IF NOT EXISTS idx_hydlogg_maskin ON hydrauliklogg(maskin_nr)")
        cols_hlog = [r[1] for r in db.execute("PRAGMA table_info(hydrauliklogg)").fetchall()]
        if 'typ' not in cols_hlog:
            db.execute("ALTER TABLE hydrauliklogg ADD COLUMN typ TEXT DEFAULT 'byte'")

init_db()

def seed_users():
    default_users = os.environ.get('DEFAULT_USERS', '')
    if not default_users:
        return
    with get_db() as db:
        count = db.execute('SELECT COUNT(*) FROM users').fetchone()[0]
        if count > 0:
            return
        for entry in default_users.split(','):
            parts = entry.strip().split(':')
            if len(parts) == 3:
                name, pin, role = parts
                db.execute('INSERT INTO users (name, pin, role) VALUES (?,?,?)', (name.strip(), pin.strip(), role.strip()))

seed_users()

def ensure_admin():
    with get_db() as db:
        exists = db.execute("SELECT id FROM users WHERE name='Felix Liljegren'").fetchone()
        if exists:
            db.execute("UPDATE users SET role='admin' WHERE name='Felix Liljegren'")
        else:
            db.execute("INSERT OR IGNORE INTO users (name, pin, role) VALUES ('Felix Liljegren','1234','admin')")

ensure_admin()

# ── Serve frontend ───────────────────────────────────────────

@app.route('/')
def index():
    return send_from_directory(BASE, 'serviceprotokoll.html')

@app.route('/manifest.json')
def manifest():
    return send_from_directory(BASE, 'manifest.json')

@app.route('/RS-logo.png')
def serve_rs_logo():
    return send_from_directory(BASE, 'RS-logo.png')

@app.route('/komp-bild.png')
def serve_komp_img():
    return send_from_directory(BASE, 'komp-bild.png')

@app.route('/balp-maskin.png')
def serve_balp_img():
    return send_from_directory(BASE, 'balp-maskin.png')

@app.route('/logo.png')
def logo():
    return send_from_directory(BASE, 'logo.png')

# ── Users & auth ─────────────────────────────────────────────

@app.route('/api/users', methods=['GET'])
def list_users():
    with get_db() as db:
        rows = db.execute('SELECT id, name, role FROM users ORDER BY name').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/users', methods=['POST'])
def create_user():
    data = request.get_json()
    name = (data.get('name') or '').strip()
    pin  = (data.get('pin') or '').strip()
    role = data.get('role', 'tekniker')
    if not name or not pin:
        return jsonify({'error': 'Namn och PIN krävs'}), 400
    with get_db() as db:
        db.execute('INSERT INTO users (name, pin, role) VALUES (?,?,?)', (name, pin, role))
    return jsonify({'ok': True}), 201

@app.route('/api/users/<int:uid>', methods=['PATCH'])
def patch_user(uid):
    data = request.get_json(force=True)
    fields = {}
    if 'pin' in data and str(data['pin']).strip():
        fields['pin'] = str(data['pin']).strip()
    if 'role' in data and data['role'] in ('admin', 'tekniker'):
        fields['role'] = data['role']
    if not fields:
        return jsonify({'error': 'Inga giltiga fält'}), 400
    set_clause = ', '.join(f'{k}=?' for k in fields)
    with get_db() as db:
        db.execute(f'UPDATE users SET {set_clause} WHERE id=?', list(fields.values()) + [uid])
    return jsonify({'ok': True})

@app.route('/api/users/<int:uid>', methods=['DELETE'])
def delete_user(uid):
    with get_db() as db:
        db.execute('DELETE FROM users WHERE id=?', (uid,))
    return jsonify({'ok': True})

@app.route('/api/login', methods=['POST'])
def login():
    data = request.get_json()
    name = (data.get('name') or '').strip()
    pin  = (data.get('pin') or '').strip()
    with get_db() as db:
        user = db.execute(
            'SELECT id, name, role FROM users WHERE name=? AND pin=?', (name, pin)
        ).fetchone()
    if not user:
        return jsonify({'error': 'Fel namn eller PIN'}), 401
    session['user'] = dict(user)
    return jsonify(dict(user))

@app.route('/api/pin-login', methods=['POST'])
def pin_login():
    pin = (request.get_json().get('pin') or '').strip()
    with get_db() as db:
        user = db.execute(
            'SELECT id, name, role FROM users WHERE pin=?', (pin,)
        ).fetchone()
    if not user:
        return jsonify({'error': 'Fel kod'}), 401
    session['user'] = dict(user)
    return jsonify({'ok': True})

@app.route('/api/logout', methods=['POST'])
def logout():
    session.pop('user', None)
    return jsonify({'ok': True})

@app.route('/api/me')
def me():
    return jsonify(session.get('user'))

# ── Auth guard ───────────────────────────────────────────────

PUBLIC_API = {'/api/login', '/api/pin-login', '/api/logout', '/api/me', '/api/users'}

@app.before_request
def require_login():
    if request.path.startswith('/api/'):
        if request.path in PUBLIC_API:
            return
        if not session.get('user'):
            return jsonify({'error': 'Ej inloggad'}), 401

# ── Machines ─────────────────────────────────────────────────

@app.route('/api/machines')
def list_machines():
    with get_db() as db:
        rows = db.execute('SELECT * FROM machines ORDER BY nr').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/machines', methods=['POST'])
def create_machine():
    data = request.get_json(force=True)
    nr = (data.get('nr') or '').strip()
    if not nr:
        # Auto-generate: LIL-YYYYMMDD-NNN (unique)
        base = 'LIL-' + datetime.now().strftime('%Y%m%d')
        with get_db() as db:
            existing = [r[0] for r in db.execute("SELECT nr FROM machines WHERE nr LIKE ?", (base + '%',)).fetchall()]
        seq = 1
        while f'{base}-{seq:03d}' in existing:
            seq += 1
        nr = f'{base}-{seq:03d}'
    with get_db() as db:
        existing = db.execute('SELECT nr FROM machines WHERE UPPER(nr)=UPPER(?)', (nr,)).fetchone()
        if existing:
            return jsonify({'error': f'Maskinnummer {nr} finns redan'}), 409
        db.execute("""INSERT INTO machines
            (nr,kund,anlaggning,fabrikat,modell,tillvnr,arsmodell,adress,stad,kontakt,telefon,agare,inkopar,beskrivning,notering,hydraulfilter,hydraulolja,hydraulslangar)
            VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
            (nr, data.get('kund',''), data.get('anlaggning',''), data.get('fabrikat',''),
             data.get('modell',''), data.get('tillvnr',''), data.get('arsmodell',''),
             data.get('adress',''), data.get('stad',''), data.get('kontakt',''),
             data.get('telefon',''), data.get('agare',''), data.get('inkopar',''),
             data.get('beskrivning',''), data.get('notering',''),
             data.get('hydraulfilter',''), data.get('hydraulolja',''), data.get('hydraulslangar','')))
    return jsonify({'nr': nr}), 201

@app.route('/api/machines/upload', methods=['POST'])
def upload_machines():
    if 'file' not in request.files:
        return jsonify({'error': 'Ingen fil'}), 400
    f = request.files['file']
    wb = openpyxl.load_workbook(f, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))

    # Find header row – stöder både Servicelista ("intern ref") och generiskt format ("maskinnummer")
    header_row = None
    servicelista_format = False
    for i, row in enumerate(rows):
        cells = [str(c or '').strip().lower() for c in row]
        if any('intern ref' in c for c in cells):
            header_row = i
            servicelista_format = True
            break
        if any('maskinnummer' in c for c in cells):
            header_row = i
            break
    if header_row is None:
        return jsonify({'error': 'Kolumn för maskinnummer saknas. Filen ska ha "Intern ref." eller "Maskinnummer".'}), 400

    hdrs_raw = [str(c or '').strip() for c in rows[header_row]]
    hdrs     = [h.lower() for h in hdrs_raw]
    def col(name): return next((i for i,h in enumerate(hdrs) if name in h), None)

    def cell(row, idx):
        if idx is None or idx >= len(row): return ''
        v = row[idx]
        if v is None: return ''
        s = str(v).strip()
        # Rensa float-format på siffror (t.ex. "106590.0" → "106590")
        if s.endswith('.0'):
            try: s = str(int(float(s)))
            except: pass
        return s

    if servicelista_format:
        cMask = col('intern ref')
        cKund = col('kundnamn')
        cAdr  = col('hämtadress rad 1')
        cSta  = col('hämtadress rad 2')
        cKont = col('kontaktperson vid problem')
        cTel  = col('telefonnummer')
        cFab  = col('märke')
        cMod  = col('typ')
        cTill = col('tillverkningsnummer')
        cArs  = col('årsmodell')
        cBesk = col('beskrivning')  # Kolumn "Beskrivning"
        cAgar = col('företag')  # Kolumn A – ägare/fakturakund

        machines = []
        for row in rows[header_row + 1:]:
            nr = cell(row, cMask)
            if not nr or nr.lower() == 'nan': continue
            adress = cell(row, cAdr)
            stad   = cell(row, cSta)
            machines.append({
                'nr':          nr,
                'kund':        cell(row, cKund),
                'anlaggning':  f"{adress}, {stad}".strip(', ') if stad else adress,
                'fabrikat':    cell(row, cFab),
                'modell':      cell(row, cMod),
                'adress':      adress,
                'stad':        stad,
                'kontakt':     cell(row, cKont),
                'telefon':     cell(row, cTel),
                'tillvnr':     cell(row, cTill),
                'arsmodell':   cell(row, cArs),
                'beskrivning': cell(row, cBesk),
                'agare':       cell(row, cAgar),
                'inkopar':     '',
                'notering':    '',
            })
    else:
        cMask = col('maskinnummer')
        cKund = col('kund')
        cAnl  = col('anläggning') or col('anlaggning')
        cFab  = col('fabrikat')
        cMod  = col('modell')
        cNot  = col('anteckn') or col('notering')
        cBesk = col('beskrivning')
        cHydf = col('hydraulfilter')
        cHydo = col('hydraulolja')
        cHyds = col('hydraulslangar')
        cGas  = col('gasdämpare') or col('gasdampare')
        cAdr  = col('adress')
        cSta  = col('stad')
        cKont = col('kontakt')
        cTel  = col('telefon')
        cTill = col('tillvnr')
        cArs  = col('årsmodell') or col('arsmodell')
        cAgar = col('ägare') or col('agare')

        machines = []
        for row in rows[header_row + 1:]:
            nr = cell(row, cMask)
            if not nr: continue
            machines.append({
                'nr':            nr,
                'kund':          cell(row, cKund),
                'anlaggning':    cell(row, cAnl),
                'fabrikat':      cell(row, cFab),
                'modell':        cell(row, cMod),
                'inkopar':       '',
                'notering':      cell(row, cNot),
                'beskrivning':   cell(row, cBesk),
                'hydraulfilter': cell(row, cHydf),
                'hydraulolja':   cell(row, cHydo),
                'hydraulslangar':cell(row, cHyds),
                'gasdampare':    cell(row, cGas),
                'adress':        cell(row, cAdr),
                'stad':          cell(row, cSta),
                'kontakt':       cell(row, cKont),
                'telefon':       cell(row, cTel),
                'tillvnr':       cell(row, cTill),
                'arsmodell':     cell(row, cArs),
                'agare':         cell(row, cAgar),
            })

    with get_db() as db:
        db.executemany("""
            INSERT INTO machines
                (nr, kund, anlaggning, fabrikat, modell, inkopar, notering, beskrivning,
                 adress, stad, kontakt, telefon, tillvnr, arsmodell, agare,
                 hydraulfilter, hydraulolja, hydraulslangar, gasdampare)
            VALUES
                (:nr, :kund, :anlaggning, :fabrikat, :modell, :inkopar, :notering, :beskrivning,
                 :adress, :stad, :kontakt, :telefon, :tillvnr, :arsmodell, :agare,
                 :hydraulfilter, :hydraulolja, :hydraulslangar, :gasdampare)
            ON CONFLICT(nr) DO UPDATE SET
                kund=excluded.kund, anlaggning=excluded.anlaggning,
                fabrikat=excluded.fabrikat, modell=excluded.modell,
                inkopar=excluded.inkopar, notering=excluded.notering,
                beskrivning=excluded.beskrivning,
                adress=excluded.adress, stad=excluded.stad,
                kontakt=excluded.kontakt, telefon=excluded.telefon,
                tillvnr=excluded.tillvnr, arsmodell=excluded.arsmodell,
                agare=excluded.agare,
                hydraulfilter=excluded.hydraulfilter,
                hydraulolja=excluded.hydraulolja,
                hydraulslangar=excluded.hydraulslangar,
                gasdampare=excluded.gasdampare
        """, machines)

    return jsonify({'imported': len(machines)})

@app.route('/api/machines/<nr>', methods=['GET'])
def get_machine(nr):
    with get_db() as db:
        row = db.execute('SELECT * FROM machines WHERE UPPER(nr)=UPPER(?)', (nr,)).fetchone()
    if not row:
        return jsonify(None), 404
    return jsonify(dict(row))

@app.route('/api/machines/<nr>', methods=['PATCH'])
def patch_machine(nr):
    data = request.get_json(force=True)
    allowed = {'mot_kaj', 'notering', 'hydraulfilter', 'hydraulolja', 'hydraulslangar', 'instruktion', 'gasdampare',
               'kund', 'anlaggning', 'fabrikat', 'modell', 'tillvnr', 'arsmodell',
               'adress', 'stad', 'kontakt', 'telefon', 'agare', 'inkopar', 'beskrivning'}
    updates = {k: v for k, v in data.items() if k in allowed}
    if not updates:
        return jsonify({'error': 'Inga giltiga fält'}), 400
    set_clause = ', '.join(f'{k}=?' for k in updates)
    with get_db() as db:
        db.execute(f'UPDATE machines SET {set_clause} WHERE UPPER(nr)=UPPER(?)',
                   list(updates.values()) + [nr])
    return jsonify({'ok': True})

@app.route('/api/machines/<nr>/history', methods=['GET'])
def get_machine_history(nr):
    with get_db() as db:
        service = db.execute(
            'SELECT id, saved_at, datum, type, tekniker, kund, modell, items_json, workflow_status FROM protocols WHERE UPPER(maskin_nr)=UPPER(?) ORDER BY saved_at DESC',
            (nr,)).fetchall()
        leverans = db.execute(
            'SELECT id, saved_at, data_json FROM leverans WHERE UPPER(data_json) LIKE ? ORDER BY saved_at DESC',
            (f'%"maskin_nr": "{nr}"%',)).fetchall()
        balp = db.execute(
            'SELECT id, saved_at, data_json FROM balp WHERE UPPER(maskin_nr)=UPPER(?) ORDER BY saved_at DESC',
            (nr,)).fetchall()
    result = []
    for r in service:
        d = dict(r)
        d['_type'] = 'service'
        d['items'] = json.loads(d.pop('items_json') or '[]')
        result.append(d)
    for r in leverans:
        d = {'_type': 'leverans', 'id': r['id'], 'saved_at': r['saved_at']}
        d.update(json.loads(r['data_json']))
        result.append(d)
    for r in balp:
        d = {'_type': 'balp', 'id': r['id'], 'saved_at': r['saved_at']}
        d.update(json.loads(r['data_json']))
        result.append(d)
    result.sort(key=lambda x: x.get('saved_at',''), reverse=True)
    return jsonify(result)

@app.route('/api/machines/<nr>/notes', methods=['GET'])
def get_machine_notes(nr):
    with get_db() as db:
        rows = db.execute('SELECT id, saved_at, author, note FROM machine_notes WHERE UPPER(maskin_nr)=UPPER(?) ORDER BY saved_at DESC', (nr,)).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/machines/<nr>/notes', methods=['POST'])
def add_machine_note(nr):
    data = request.get_json()
    with get_db() as db:
        db.execute('INSERT INTO machine_notes (maskin_nr, saved_at, author, note) VALUES (?,?,?,?)',
                   (nr, datetime.now().isoformat(), data.get('author',''), data.get('note','')))
    return jsonify({'ok': True})

# ── Leveransgodkännande ──────────────────────────────────────

def init_leverans_db():
    with get_db() as db:
        db.execute('''CREATE TABLE IF NOT EXISTS leverans (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            saved_at  TEXT,
            tekniker  TEXT,
            data_json TEXT
        )''')

init_leverans_db()

def init_balp_db():
    with get_db() as db:
        db.execute('''CREATE TABLE IF NOT EXISTS balp (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            saved_at  TEXT,
            tekniker  TEXT,
            maskin_nr TEXT,
            data_json TEXT
        )''')

init_balp_db()

@app.route('/api/balp', methods=['POST'])
def save_balp():
    data = request.get_json()
    with get_db() as db:
        db.execute('INSERT INTO balp (saved_at, tekniker, maskin_nr, data_json) VALUES (?,?,?,?)',
                   (datetime.now().isoformat(), data.get('tekniker',''), data.get('maskin_nr',''), json.dumps(data, ensure_ascii=False)))
    return jsonify({'ok': True})

@app.route('/api/balp', methods=['GET'])
def list_balp():
    with get_db() as db:
        rows = db.execute('SELECT id, saved_at, tekniker, maskin_nr, data_json FROM balp ORDER BY saved_at DESC').fetchall()
    return jsonify([{**dict(r), 'data': json.loads(r['data_json'])} for r in rows])

@app.route('/api/balp/<int:balp_id>', methods=['DELETE'])
def delete_balp(balp_id):
    with get_db() as db:
        db.execute('DELETE FROM balp WHERE id=?', (balp_id,))
    return jsonify({'ok': True})

# ── Fordonsregister & Fordonskontroll ───────────────────────────────────────

FORDON_DB = os.path.join(DATA_DIR, 'fordon-register.db')

def get_fordon_db():
    conn = sqlite3.connect(FORDON_DB)
    conn.row_factory = sqlite3.Row
    return conn

def init_fordon_db():
    with get_fordon_db() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS fordon (
            regnr      TEXT PRIMARY KEY,
            typ        TEXT DEFAULT '',
            marke      TEXT DEFAULT '',
            modell     TEXT DEFAULT '',
            arsmodell  TEXT DEFAULT '',
            besiktning TEXT DEFAULT '',
            forare     TEXT DEFAULT '',
            avdelning  TEXT DEFAULT '',
            notering   TEXT DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS fordonskontroll (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            regnr     TEXT NOT NULL,
            datum     TEXT,
            tekniker  TEXT,
            saved_at  TEXT,
            data_json TEXT
        );
        CREATE TABLE IF NOT EXISTS materialkontroll (
            id        INTEGER PRIMARY KEY AUTOINCREMENT,
            regnr     TEXT NOT NULL,
            datum     TEXT,
            ansvarig  TEXT,
            saved_at  TEXT,
            data_json TEXT
        );
        """)

init_fordon_db()

@app.route('/api/fordon/upload', methods=['POST'])
def upload_fordon():
    f = request.files.get('file')
    if not f:
        return jsonify({'error': 'Ingen fil'}), 400
    path = os.path.join(UPLOADS, 'fordon-register.xlsx')
    f.save(path)
    wb = openpyxl.load_workbook(path, data_only=True)
    imported = 0
    with get_fordon_db() as db:
        db.execute('DELETE FROM fordon')
        # En flik: "Maskin och fordonsregister" — ID/Regnr, Typ, Märke, Modell, Årsmodell, Besiktning/Serienr, Förare/Ansvarig, Placering/Avdelning, Notering
        sheet_name = next((s for s in wb.sheetnames if 'maskin' in s.lower() or 'fordon' in s.lower()), wb.sheetnames[0])
        for row in wb[sheet_name].iter_rows(min_row=3, values_only=True):
            regnr = str(row[0]).strip() if row[0] else None
            if not regnr or regnr.lower().startswith('fyll'): continue
            db.execute('INSERT OR REPLACE INTO fordon (regnr,typ,marke,modell,arsmodell,besiktning,forare,avdelning,notering) VALUES (?,?,?,?,?,?,?,?,?)',
                (regnr, str(row[1] or ''), str(row[2] or ''), str(row[3] or ''),
                 str(row[4] or ''), str(row[5] or ''), str(row[6] or ''),
                 str(row[7] or ''), str(row[8] or '')))
            imported += 1
    return jsonify({'imported': imported})

@app.route('/api/fordon', methods=['GET'])
def list_fordon():
    with get_fordon_db() as db:
        rows = db.execute('SELECT * FROM fordon ORDER BY regnr').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/fordonskontroll', methods=['POST'])
def save_fordonskontroll():
    data = request.get_json()
    with get_fordon_db() as db:
        db.execute('INSERT INTO fordonskontroll (regnr, datum, tekniker, saved_at, data_json) VALUES (?,?,?,?,?)',
                   (data.get('regnr',''), data.get('datum',''), data.get('tekniker',''),
                    datetime.now().isoformat(), json.dumps(data, ensure_ascii=False)))
    return jsonify({'ok': True})

@app.route('/api/fordonskontroll', methods=['GET'])
def list_fordonskontroll():
    with get_fordon_db() as db:
        rows = db.execute('SELECT id, saved_at, regnr, tekniker, datum, data_json FROM fordonskontroll ORDER BY saved_at DESC').fetchall()
    return jsonify([{**dict(r), 'data': json.loads(r['data_json'])} for r in rows])

@app.route('/api/fordonskontroll/export', methods=['GET'])
def export_fordonskontroll():
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    with get_fordon_db() as db:
        rows = db.execute('SELECT id, saved_at, regnr, tekniker, datum, data_json FROM fordonskontroll ORDER BY saved_at DESC').fetchall()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Fordonskontroller'

    HEADER_FILL = PatternFill('solid', fgColor='1A3A5C')
    HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=11)
    BODY_FONT   = Font(name='Arial', size=10)
    THIN = Side(style='thin', color='CFD8DC')
    BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    OK_FILL   = PatternFill('solid', fgColor='C8E6C9')
    BEV_FILL  = PatternFill('solid', fgColor='FFF9C4')
    AKUT_FILL = PatternFill('solid', fgColor='FFCDD2')

    headers = ['Protokoll-ID', 'Datum', 'Regnr', 'Förare', 'Km-ställning', 'Nr', 'Kontrollpunkt', 'Status', 'Anmärkning', 'Övrigt']
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[1].height = 30

    row_idx = 2
    for r in rows:
        d = json.loads(r['data_json'])
        datum   = d.get('datum', r['datum'] or '')
        regnr   = d.get('regnr', r['regnr'] or '')
        tekniker= d.get('tekniker', r['tekniker'] or '')
        km      = d.get('km', '')
        ovrig   = d.get('summary_notes', '')
        items   = d.get('items', [])

        if not items:
            # Skriv en rad ändå
            for col, val in enumerate([r['id'], datum, regnr, tekniker, km, '', '', '', '', ovrig], 1):
                c = ws.cell(row=row_idx, column=col, value=val)
                c.font = BODY_FONT; c.border = BORDER; c.alignment = Alignment(vertical='center')
            ws.row_dimensions[row_idx].height = 18
            row_idx += 1
        else:
            for i, item in enumerate(items):
                val_text = {'1': 'OK', '2': 'Bevaka', '3': 'Akut'}.get(item.get('value', ''), '')
                fill = {'OK': OK_FILL, 'Bevaka': BEV_FILL, 'Akut': AKUT_FILL}.get(val_text)
                for col, val in enumerate([
                    r['id'] if i == 0 else '',
                    datum   if i == 0 else '',
                    regnr   if i == 0 else '',
                    tekniker if i == 0 else '',
                    km      if i == 0 else '',
                    item.get('nr',''), item.get('label',''), val_text, item.get('comment',''),
                    ovrig   if i == 0 else ''
                ], 1):
                    c = ws.cell(row=row_idx, column=col, value=val)
                    c.font = BODY_FONT; c.border = BORDER
                    c.alignment = Alignment(vertical='center', wrap_text=(col in (7,9)))
                    if fill and col == 8:
                        c.fill = fill
                ws.row_dimensions[row_idx].height = 18
                row_idx += 1

    col_widths = [12, 12, 12, 18, 10, 6, 32, 10, 28, 28]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    ws.freeze_panes = 'A2'

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"fordonskontroller_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')

@app.route('/api/fordonskontroll/<int:fk_id>', methods=['DELETE'])
def delete_fordonskontroll(fk_id):
    with get_fordon_db() as db:
        db.execute('DELETE FROM fordonskontroll WHERE id=?', (fk_id,))
    return jsonify({'ok': True})

@app.route('/api/materialkontroll', methods=['POST'])
def save_materialkontroll():
    data = request.get_json()
    with get_fordon_db() as db:
        db.execute('INSERT INTO materialkontroll (regnr, datum, ansvarig, saved_at, data_json) VALUES (?,?,?,?,?)',
                   (data.get('regnr',''), data.get('datum',''), data.get('ansvarig',''),
                    datetime.now().isoformat(), json.dumps(data, ensure_ascii=False)))
    return jsonify({'ok': True})

@app.route('/api/materialkontroll', methods=['GET'])
def list_materialkontroll():
    with get_fordon_db() as db:
        rows = db.execute('SELECT id, saved_at, regnr, ansvarig, datum, data_json FROM materialkontroll ORDER BY saved_at DESC').fetchall()
    return jsonify([{**dict(r), 'data': json.loads(r['data_json'])} for r in rows])

@app.route('/api/installationsprotokoll', methods=['POST'])
def save_installationsprotokoll():
    data = request.get_json()
    with get_db() as db:
        db.execute('''CREATE TABLE IF NOT EXISTS installationsprotokoll (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            saved_at TEXT, montör TEXT, kund TEXT, datum TEXT, data_json TEXT)''')
        db.execute('INSERT INTO installationsprotokoll (saved_at, montör, kund, datum, data_json) VALUES (?,?,?,?,?)',
                   (datetime.now().isoformat(), data.get('montör',''), data.get('kund',''),
                    data.get('datum',''), json.dumps(data, ensure_ascii=False)))
    return jsonify({'ok': True})

@app.route('/api/installationsprotokoll', methods=['GET'])
def list_installationsprotokoll():
    with get_db() as db:
        db.execute('''CREATE TABLE IF NOT EXISTS installationsprotokoll (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            saved_at TEXT, montör TEXT, kund TEXT, datum TEXT, data_json TEXT)''')
        rows = db.execute('SELECT id, saved_at, montör, kund, datum FROM installationsprotokoll ORDER BY saved_at DESC').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/installationsprotokoll/<int:rec_id>', methods=['DELETE'])
def delete_installationsprotokoll(rec_id):
    with get_db() as db:
        db.execute('DELETE FROM installationsprotokoll WHERE id=?', (rec_id,))
    return jsonify({'ok': True})

@app.route('/api/leverans', methods=['POST'])
def save_leverans():
    data = request.get_json()
    with get_db() as db:
        db.execute('INSERT INTO leverans (saved_at, tekniker, data_json) VALUES (?,?,?)',
                   (datetime.now().isoformat(), data.get('tekniker',''), json.dumps(data, ensure_ascii=False)))
    return jsonify({'ok': True})

@app.route('/api/leverans', methods=['GET'])
def list_leverans():
    with get_db() as db:
        rows = db.execute('SELECT id, saved_at, tekniker, data_json FROM leverans ORDER BY saved_at DESC').fetchall()
    return jsonify([{**dict(r), 'data': json.loads(r['data_json'])} for r in rows])

# ── Protocols ────────────────────────────────────────────────

@app.route('/api/protocols/search')
def search_protocols():
    limit  = min(int(request.args.get('limit', 100)), 500)
    kund   = request.args.get('kund', '').strip()
    from_d = request.args.get('from', '').strip()
    to_d   = request.args.get('to', '').strip()

    query  = 'SELECT * FROM protocols WHERE 1=1'
    params = []
    if kund:
        query += ' AND LOWER(kund) = LOWER(?)'
        params.append(kund)
    if from_d:
        query += ' AND datum >= ?'
        params.append(from_d)
    if to_d:
        query += ' AND datum <= ?'
        params.append(to_d)
    query += ' ORDER BY datum DESC, saved_at DESC LIMIT ?'
    params.append(limit)

    with get_db() as db:
        rows = db.execute(query, params).fetchall()
    result = []
    for r in rows:
        d = dict(r)
        d['items']          = json.loads(d['items_json'] or '[]')
        d['resolved']       = json.loads(d.get('resolved_json') or '{}')
        d['workflow_log']   = json.loads(d.get('workflow_log') or '[]')
        if not d.get('workflow_status'): d['workflow_status'] = 'inkommen'
        del d['items_json'], d['resolved_json']
        result.append(d)
    return jsonify(result)

@app.route('/api/protocols/recent')
def recent_protocols():
    limit = min(int(request.args.get('limit', 15)), 50)
    with get_db() as db:
        rows = db.execute(
            'SELECT * FROM protocols ORDER BY saved_at DESC LIMIT ?', (limit,)
        ).fetchall()
    result = []
    for r in rows:
        d = dict(r)
        d['items']    = json.loads(d['items_json'] or '[]')
        d['resolved']      = json.loads(d.get('resolved_json') or '{}')
        d['workflow_log']  = json.loads(d.get('workflow_log') or '[]')
        if 'workflow_status' not in d or not d['workflow_status']:
            d['workflow_status'] = 'inkommen'
        del d['items_json'], d['resolved_json']
        result.append(d)
    return jsonify(result)

@app.route('/api/protocols/<maskin_nr>')
def get_protocols(maskin_nr):
    with get_db() as db:
        rows = db.execute(
            'SELECT * FROM protocols WHERE UPPER(maskin_nr)=UPPER(?) ORDER BY saved_at DESC LIMIT 30',
            (maskin_nr,)
        ).fetchall()
    result = []
    for r in rows:
        d = dict(r)
        d['items']          = json.loads(d['items_json'] or '[]')
        d['resolved']       = json.loads(d.get('resolved_json') or '{}')
        d['workflow_log']   = json.loads(d.get('workflow_log') or '[]')
        if not d.get('workflow_status'):
            d['workflow_status'] = 'inkommen'
        del d['items_json'], d['resolved_json']
        result.append(d)
    return jsonify(result)

@app.route('/api/protocols/<int:protocol_id>', methods=['DELETE'])
def delete_protocol(protocol_id):
    with get_db() as db:
        db.execute('DELETE FROM protocols WHERE id=?', (protocol_id,))
    return jsonify({'ok': True})

@app.route('/api/leverans/<int:leverans_id>', methods=['DELETE'])
def delete_leverans(leverans_id):
    with get_db() as db:
        db.execute('DELETE FROM leverans WHERE id=?', (leverans_id,))
    return jsonify({'ok': True})

@app.route('/api/protocols/<int:protocol_id>/resolve', methods=['PATCH'])
def resolve_item(protocol_id):
    data = request.get_json()
    nr   = data.get('nr')
    if not nr:
        return jsonify({'error': 'nr krävs'}), 400
    with get_db() as db:
        row = db.execute('SELECT resolved_json FROM protocols WHERE id=?', (protocol_id,)).fetchone()
        if not row:
            return jsonify({'error': 'Protokoll ej hittat'}), 404
        resolved = json.loads(row['resolved_json'] or '{}')
        if data.get('undo'):
            resolved.pop(nr, None)
        else:
            resolved[nr] = {
                'date': datetime.now().strftime('%Y-%m-%d'),
                'by':   data.get('by', '').strip()
            }
        db.execute('UPDATE protocols SET resolved_json=? WHERE id=?',
                   (json.dumps(resolved, ensure_ascii=False), protocol_id))
    return jsonify({'resolved': resolved})

# ── Workflow ──────────────────────────────────────────────────

WORKFLOW_STEPS = ['inkommen', 'hanterad']
WORKFLOW_LABELS = {
    'inkommen': 'Inkommen',
    'hanterad': 'Hanterad',
}

@app.route('/api/protocols/<int:protocol_id>/workflow', methods=['PATCH'])
def update_workflow(protocol_id):
    data   = request.get_json()
    status = data.get('status')
    note   = data.get('note', '').strip()
    user   = session.get('user', {})
    if status not in WORKFLOW_STEPS:
        return jsonify({'error': 'Ogiltigt status'}), 400
    with get_db() as db:
        row = db.execute('SELECT workflow_log FROM protocols WHERE id=?', (protocol_id,)).fetchone()
        if not row:
            return jsonify({'error': 'Protokoll ej hittat'}), 404
        log = json.loads(row['workflow_log'] or '[]')
        log.append({
            'status': status,
            'label':  WORKFLOW_LABELS[status],
            'by':     user.get('name', 'Okänd'),
            'at':     datetime.now().strftime('%Y-%m-%d %H:%M'),
            'note':   note,
        })
        db.execute('UPDATE protocols SET workflow_status=?, workflow_log=? WHERE id=?',
                   (status, json.dumps(log, ensure_ascii=False), protocol_id))
    return jsonify({'workflow_status': status, 'workflow_log': log})

@app.route('/api/protocols', methods=['POST'])
def save_protocol():
    data = request.get_json()
    if not data or not data.get('maskin_nr'):
        return jsonify({'error': 'maskin_nr krävs'}), 400
    with get_db() as db:
        db.execute("""
            INSERT INTO protocols (maskin_nr, datum, type, machine_type, kund, anlaggning, tekniker, modell, items_json, saved_at)
            VALUES (?,?,?,?,?,?,?,?,?,?)
        """, (
            data['maskin_nr'].upper(),
            data.get('datum', ''),
            data.get('type', ''),
            data.get('machine_type', ''),
            data.get('kund', ''),
            data.get('anlaggning', ''),
            data.get('tekniker', ''),
            data.get('modell', ''),
            json.dumps(data.get('items', []), ensure_ascii=False),
            datetime.now().isoformat()
        ))
        pid = db.execute('SELECT last_insert_rowid()').fetchone()[0]
    return jsonify({'id': pid}), 201

# ── Photo upload ─────────────────────────────────────────────

ALLOWED_EXT = {'.jpg', '.jpeg', '.png', '.gif', '.webp', '.heic'}

@app.route('/api/photos', methods=['POST'])
def upload_photo():
    if 'photo' not in request.files:
        return jsonify({'error': 'Ingen fil'}), 400
    f = request.files['photo']
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ALLOWED_EXT:
        return jsonify({'error': 'Otillåtet filformat'}), 400
    filename = uuid.uuid4().hex + ext
    f.save(os.path.join(UPLOADS, filename))
    return jsonify({'filename': filename}), 201

@app.route('/uploads/<filename>')
def serve_upload(filename):
    return send_from_directory(UPLOADS, filename)

@app.route('/api/machines/export', methods=['GET'])
def export_machines():
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    import io as _io
    with get_db() as db:
        rows = db.execute('SELECT * FROM machines ORDER BY nr').fetchall()
    wb = Workbook()
    ws = wb.active
    ws.title = 'Maskinregister'
    headers = ['Maskinnummer','Kund','Anläggning','Fabrikat','Modell','Tillvnr','Årsmodell',
               'Adress','Stad','Kontakt','Telefon','Ägare','Beskrivning','Notering',
               'Hydraulfilter','Hydraulolja','Hydraulslangar','Gasdämpare']
    cols    = ['nr','kund','anlaggning','fabrikat','modell','tillvnr','arsmodell',
               'adress','stad','kontakt','telefon','agare','beskrivning','notering',
               'hydraulfilter','hydraulolja','hydraulslangar','gasdampare']
    header_fill = PatternFill('solid', fgColor='1A3A5C')
    header_font = Font(bold=True, color='FFFFFF', name='Arial', size=10)
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.fill  = header_fill
        cell.font  = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    for ri, row in enumerate(rows, 2):
        d = dict(row)
        for ci, col in enumerate(cols, 1):
            ws.cell(row=ri, column=ci, value=d.get(col) or '')
    col_widths = [18,30,35,20,20,18,12,30,20,25,18,25,35,35,25,25,25,25]
    for ci, w in enumerate(col_widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=ci).column_letter].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    from flask import Response
    fname = f"maskinregister_{datetime.now().strftime('%Y%m%d')}.xlsx"
    return Response(buf.read(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename="{fname}"'})

@app.route('/api/machines/<maskin_nr>/typskylt', methods=['POST'])
def upload_typskylt(maskin_nr):
    if 'photo' not in request.files:
        return jsonify({'error': 'Ingen fil'}), 400
    f = request.files['photo']
    ext = os.path.splitext(f.filename)[1].lower()
    if ext not in ALLOWED_EXT:
        return jsonify({'error': 'Otillåtet filformat'}), 400
    safe_nr = ''.join(c for c in maskin_nr if c.isalnum() or c in '-_')
    filename = f'typskylt_{safe_nr}{ext}'
    f.save(os.path.join(UPLOADS, filename))
    with get_db() as db:
        db.execute("UPDATE machines SET typskylt=? WHERE UPPER(nr)=UPPER(?)", (filename, maskin_nr))
    return jsonify({'filename': filename}), 201

@app.route('/api/machines/<maskin_nr>/typskylt', methods=['GET'])
def get_typskylt(maskin_nr):
    with get_db() as db:
        row = db.execute("SELECT typskylt FROM machines WHERE UPPER(nr)=UPPER(?)", (maskin_nr,)).fetchone()
    if not row or not row['typskylt']:
        return jsonify({'filename': None})
    return jsonify({'filename': row['typskylt']})

@app.route('/api/hydrauliklogg/<maskin_nr>', methods=['GET'])
def get_hydrauliklogg(maskin_nr):
    with get_db() as db:
        rows = db.execute(
            'SELECT * FROM hydrauliklogg WHERE UPPER(maskin_nr)=UPPER(?) ORDER BY datum DESC',
            [maskin_nr]).fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/hydrauliklogg', methods=['POST'])
def post_hydrauliklogg():
    d = request.get_json(force=True)
    with get_db() as db:
        db.execute(
            'INSERT INTO hydrauliklogg (maskin_nr,datum,filter,olja,slangar,notering,saved_at,typ) VALUES (?,?,?,?,?,?,?,?)',
            [d.get('maskin_nr',''), d.get('datum',''), int(d.get('filter',0)),
             int(d.get('olja',0)), int(d.get('slangar',0)), d.get('notering',''),
             datetime.now().isoformat(), d.get('typ','byte')])
    return jsonify({'ok': True})

@app.route('/api/hydrauliklogg/<int:log_id>', methods=['DELETE'])
def delete_hydrauliklogg(log_id):
    with get_db() as db:
        db.execute('DELETE FROM hydrauliklogg WHERE id=?', [log_id])
    return jsonify({'ok': True})

@app.route('/api/offers', methods=['POST'])
def create_offer():
    d = request.get_json(force=True)
    with get_db() as db:
        cur = db.execute(
            'INSERT INTO offers (protocol_id,maskin_nr,kund,anlaggning,modell,datum,tekniker,items_json,saved_at,status) VALUES (?,?,?,?,?,?,?,?,?,?)',
            [d.get('protocol_id'), d.get('maskin_nr',''), d.get('kund',''), d.get('anlaggning',''),
             d.get('modell',''), d.get('datum',''), d.get('tekniker',''),
             json.dumps(d.get('items', []), ensure_ascii=False),
             datetime.now().isoformat(), d.get('status','utkast')])
        offer_id = cur.lastrowid
    return jsonify({'id': offer_id}), 201

@app.route('/api/offers', methods=['GET'])
def list_offers():
    with get_db() as db:
        rows = db.execute('SELECT id,protocol_id,maskin_nr,kund,datum,tekniker,status,saved_at FROM offers ORDER BY saved_at DESC').fetchall()
    return jsonify([dict(r) for r in rows])

@app.route('/api/offers/<int:offer_id>', methods=['GET'])
def get_offer(offer_id):
    with get_db() as db:
        row = db.execute('SELECT * FROM offers WHERE id=?', [offer_id]).fetchone()
    if not row: return jsonify(None), 404
    d = dict(row)
    d['items'] = json.loads(d.pop('items_json') or '[]')
    return jsonify(d)

@app.route('/api/offers/<int:offer_id>', methods=['PATCH'])
def patch_offer(offer_id):
    d = request.get_json(force=True)
    allowed = {'status', 'items_json'}
    if 'items' in d:
        d['items_json'] = json.dumps(d.pop('items'), ensure_ascii=False)
    updates = {k: v for k, v in d.items() if k in allowed}
    if not updates: return jsonify({'error': 'Inga fält'}), 400
    set_clause = ', '.join(f'{k}=?' for k in updates)
    with get_db() as db:
        db.execute(f'UPDATE offers SET {set_clause} WHERE id=?', list(updates.values()) + [offer_id])
    return jsonify({'ok': True})

@app.route('/api/offers/<int:offer_id>', methods=['DELETE'])
def delete_offer(offer_id):
    with get_db() as db:
        db.execute('DELETE FROM offers WHERE id=?', [offer_id])
    return jsonify({'ok': True})

@app.route('/spiris-test.html')
def serve_spiris_test():
    return send_from_directory(BASE, 'spiris-test.html')

# ── Spiris proxy ─────────────────────────────────────────────────────────────
SPIRIS_HOST = 'publicapi.tid.vismaspcs.se'

def spiris_token():
    app_id = os.environ.get('SPIRIS_APP_ID', '')
    secret = os.environ.get('SPIRIS_SECRET', '')
    if not app_id or not secret:
        return None, 'SPIRIS_APP_ID eller SPIRIS_SECRET saknas i miljövariabler'
    cred = base64.b64encode(f'{app_id}:{secret}'.encode()).decode()
    req = urllib.request.Request(
        f'https://{SPIRIS_HOST}/v1/Auth/Token',
        method='POST',
        headers={'Authorization': f'Basic {cred}', 'Content-Type': 'application/json', 'Content-Length': '0'}
    )
    try:
        with urllib.request.urlopen(req, timeout=10) as r:
            data = json.loads(r.read())
            return data.get('accessToken'), None
    except urllib.error.HTTPError as e:
        body = e.read().decode('utf-8', errors='replace')
        return None, f'HTTP {e.code}: {body}'
    except Exception as e:
        return None, str(e)

def spiris_get(token, path):
    req = urllib.request.Request(
        f'https://{SPIRIS_HOST}{path}',
        headers={'Authorization': f'Bearer {token}', 'Accept': 'application/json'}
    )
    with urllib.request.urlopen(req, timeout=10) as r:
        return json.loads(r.read())

@app.route('/api/spiris/projekt/<projekt_nr>')
def spiris_projekt(projekt_nr):
    token, err = spiris_token()
    if err:
        return jsonify({'error': err}), 500
    try:
        projekt = spiris_get(token, f'/v1/Core/Projects?page=1&filter.query={urllib.parse.quote(projekt_nr)}')
        return jsonify(projekt)
    except Exception as e:
        return jsonify({'error': str(e)}), 500

@app.route('/api/spiris/test')
def spiris_test():
    token, err = spiris_token()
    if err:
        return jsonify({'ok': False, 'error': err}), 500
    return jsonify({'ok': True, 'token_length': len(token)})

if __name__ == '__main__':
    import socket
    host = '0.0.0.0'
    port = 5000
    local_ip = socket.gethostbyname(socket.gethostname())
    print(f'\n  Serviceprotokoll-server startad!')
    print(f'  Öppna i webbläsare: http://localhost:{port}')
    print(f'  Från andra enheter: http://{local_ip}:{port}')
    print(f'  Tryck Ctrl+C för att stänga av\n')
    app.run(host=host, port=port, debug=False)
